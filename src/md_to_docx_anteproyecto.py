"""
md_to_docx_anteproyecto.py
============================
Convierte el anteproyecto v4 de markdown a .docx,
respetando el formato del original (portada, secciones, tablas, referencias).

Uso: python3 md_to_docx_anteproyecto.py <input.md> <output.docx>
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_borders(cell):
    """Add borders to a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for border_name in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "000000")
        tcBorders.append(border)
    tcPr.append(tcBorders)


def parse_inline(text):
    """Parse bold (**text**) and italic (*text*) into runs."""
    parts = []
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        token = m.group(0)
        if token.startswith("**"):
            parts.append((token[2:-2], True))
        else:
            parts.append((token[1:-1], "italic"))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    return parts


def add_runs(paragraph, text, base_size=11):
    """Add styled runs to a paragraph."""
    for content, style in parse_inline(text):
        run = paragraph.add_run(content)
        run.font.size = Pt(base_size)
        if style is True:
            run.bold = True
        elif style == "italic":
            run.italic = True


def add_heading_para(doc, text, level, base_size=11):
    """Add a heading paragraph (1.1, 1.2, etc.)."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(base_size)
    run.bold = True
    return p


def add_body_para(doc, text, base_size=11):
    """Add a body paragraph (wrapped lines)."""
    lines = text.split("\n")
    for line in lines:
        p = doc.add_paragraph()
        add_runs(p, line, base_size)
    return p


def add_unordered_list(doc, items, base_size=11):
    """Add an unordered list."""
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, item, base_size)


def add_ordered_list(doc, items, base_size=11):
    """Add an ordered list."""
    for item in items:
        p = doc.add_paragraph(style="List Number")
        add_runs(p, item, base_size)


def add_table(doc, rows, base_size=11):
    """Add a table with given rows (list of lists)."""
    if not rows:
        return
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            if j < len(row):
                cell = table.rows[i].cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                add_runs(p, cell_text, base_size)


def parse_table_block(lines, start_idx):
    """Parse a markdown table block."""
    table_rows = []
    i = start_idx
    while i < len(lines) and "|" in lines[i]:
        line = lines[i].strip()
        if re.match(r"^\|?[\s\-|:]+\|?$", line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        table_rows.append(cells)
        i += 1
    return table_rows, i


def add_image_placeholder(doc, alt_text):
    """Add a placeholder paragraph for an image (we don't embed images)."""
    p = doc.add_paragraph()
    run = p.add_run(f"[Imagen: {alt_text}]")
    run.italic = True
    run.font.color.rgb = None
    return p


def convert_md_to_docx(md_path, docx_path):
    """Convert a markdown file to .docx."""
    md_text = Path(md_path).read_text(encoding="utf-8")
    lines = md_text.split("\n")

    doc = Document()

    # Set default style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # ---- PORTADA ----
    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Universidad Nacional de Asunción")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Facultad Politécnica")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Ingeniería en Informática")
    r.italic = True
    r.font.size = Pt(12)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Propuesta de Proyecto Final de Carrera")
    r.italic = True
    r.font.size = Pt(12)

    doc.add_paragraph()
    doc.add_paragraph()

    # Extract title from the markdown (first heading)
    title = ""
    for line in lines:
        if line.startswith("# ") and not line.startswith("## "):
            title = line[2:].strip()
            break

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)

    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Julio 2026")
    r.italic = True
    r.font.size = Pt(12)

    doc.add_page_break()

    i = 0
    skip_until = -1

    # Skip until we find the first main section "1. SÍNTESIS"
    while i < len(lines) and not lines[i].startswith("1. SÍNTESIS"):
        i += 1

    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        # Skip lines that are part of the title
        if i < 5:
            i += 1
            continue

        # Detect headings
        if re.match(r"^\d+\.\s+[A-ZÁÉÍÓÚÑ]", line):
            # Level 1 heading (e.g. "1. SÍNTESIS DEL PROYECTO")
            add_heading_para(doc, line, 1, base_size=13)
        elif re.match(r"^\d+\.\d+\s+", line):
            # Level 2 heading (e.g. "1.1 TÍTULO INICIAL DEL PROYECTO")
            add_heading_para(doc, line, 2, base_size=12)
        elif re.match(r"^\d+\.\d+\.\d+\s+", line):
            # Level 3 heading
            add_heading_para(doc, line, 3, base_size=11)
        elif line.startswith("SINTETIZAR EL PROYECTO"):
            p = doc.add_paragraph()
            r = p.add_run("SINTETIZAR EL PROYECTO EN TRES PALABRAS CLAVE")
            r.bold = True
            r.font.size = Pt(11)
        # Detect unordered list
        elif re.match(r"^[-*]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].rstrip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].rstrip()))
                i += 1
            add_unordered_list(doc, items)
            continue
        # Detect ordered list
        elif re.match(r"^\d+[.)]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\d+[.)]\s+", lines[i].rstrip()):
                items.append(re.sub(r"^\d+[.)]\s+", "", lines[i].rstrip()))
                i += 1
            add_ordered_list(doc, items)
            continue
        # Detect table
        elif "|" in line and i + 1 < len(lines) and "|" in lines[i + 1]:
            # Table detected
            table_rows, i = parse_table_block(lines, i)
            if table_rows:
                add_table(doc, table_rows)
            continue
        # Regular paragraph
        else:
            add_body_para(doc, line)

        i += 1

    doc.save(docx_path)
    print(f"OK Saved: {docx_path} ({len(md_text.split(chr(10)))} lines processed)")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python3 md_to_docx_anteproyecto.py <input.md> <output.docx>")
        sys.exit(1)
    convert_md_to_docx(sys.argv[1], sys.argv[2])
